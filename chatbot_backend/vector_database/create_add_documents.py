import os
from datetime import datetime
from dotenv import load_dotenv
from uuid import uuid4
from pathlib import Path
from openai import OpenAI
from pydantic import BaseModel, Field
from tqdm import tqdm
from litellm import completion
from docx import Document as DocxDocument
from langchain_core.documents import Document

from chatbot_backend.vector_database.documents_metadata import doc_metadata
from chatbot_backend.all_models.embedding_model import vector_store

load_dotenv(override=True)

MODEL = "gpt-4.1-nano"
EMBEDDING_MODEL = "BAAI/bge-m3"
AVERAGE_CHUNK_SIZE = 500
BASE_DIR = Path(__file__).resolve().parent

KNOWLEDGE_BASE_PATH = BASE_DIR / "mit_vn_documents_level"
"""
The directory specified by `KNOWLEDGE_BASE_PATH` contains subdirectories, each representing a specific document category. 
Each subdirectory holds the corresponding `.docx` files for that category. 
Note that the specific document files are hidden here; please replace them with your own documents. 
Similarly, the category subdirectories can also be customized as needed.
"""

openai = OpenAI()

class Result(BaseModel):
    page_content: str
    metadata: dict
    id:str

class Chunk(BaseModel):
    headline: str = Field(description="A brief heading for this chunk, typically a few words, that is most likely to be surfaced in a query")
    summary: str = Field(description="A few sentences summarizing the content of this chunk to answer common questions")
    original_text: str = Field(description="The original text of this chunk from the provided document, exactly as is, not changed in any way")

    def as_result(self, document):
        metadata = {
            "source": document['source'],
            "file_name": document['file_name'],
            "type": document['type'],
            "last_update": datetime.now().strftime('%d/%m/%Y %H:%M:%S'),
            "doc_id": str(uuid4())
        }
        return Result(
           page_content=self.headline + "\n\n" + self.summary + "\n\n" + self.original_text,
           metadata=metadata,
           id=str(uuid4())
        )

class Chunks(BaseModel):
    chunks:list[Chunk]

def fetch_document():
  documents = []

  for folder in KNOWLEDGE_BASE_PATH.iterdir():
    if not folder.is_dir():
      continue

    doc_type = folder.name

    for file_path in folder.rglob("*.docx"):
      file_name = file_path.name
      source = ""
      for metadata in doc_metadata:
        if metadata["doc_name"] == file_path.name:
          file_name = metadata["title"]
          source = metadata["source_url"]
          break

      try:
        doc = DocxDocument(file_path)
        text = "\n".join(
            [para.text for para in doc.paragraphs if para.text.strip()]
        )

        document_metadata = {
            "type": doc_type,
            "source": source,
            "file_name": file_name,
            "text": text,
        }
        documents.append(document_metadata)

      except Exception as e:
        print(f"Error when reading file {file_path.name}: {e}")

  print(f"Loaded {len(documents)} documents")
  return documents

documents = fetch_document()

def make_prompt(document):
    how_many = (len(document["text"]) // AVERAGE_CHUNK_SIZE) + 1
    return f"""
You take a document and you split the document into overlapping chunks for a KnowledgeBase.

The document belongs to MIT University VN (Trường Đại học Công nghệ miền Đông).
The document is of type: {document["type"]}
This document originates from the URL link: {document["source"]}

A chatbot will use these chunks to answer questions about the university.
You should divide up the document as you see fit, being sure that the entire document is returned in the chunks - don't leave anything out.
This document should probably be split into {how_many} chunks, but you can have more or less as appropriate.
There should be overlap between the chunks as appropriate; typically about 25% overlap or about 50 words, so you have the same text in multiple chunks for best retrieval results.

For each chunk, you should provide a headline, a summary, and the original text of the chunk.
Together your chunks should represent the entire document with overlap.

Requirements:
- Please write the response (headlines, summaries, and the original text of the chunk.) in Vietnamese.
- Format the final output clearly using Markdown (e.g., using headings, bullet points, or blockquotes for each chunk).

Here is the document:

{document["text"]}

Respond with the chunks.
"""

def make_messages(document):
    return [
        {"role": "user", "content": make_prompt(document)},
    ]

def process_document(document):
    messages = make_messages(document)
    response = completion(model=MODEL, messages=messages, response_format=Chunks)
    reply = response.choices[0].message.content
    doc_as_chunks = Chunks.model_validate_json(reply).chunks
    return [chunk.as_result(document) for chunk in doc_as_chunks]

def create_chunks(documents):
  chunks = []
  for doc in tqdm(documents):
    results = process_document(doc)

    lc_documents = [
        Document(
            page_content=res.page_content, metadata=res.metadata, id=res.id
        )
        for res in results
    ]

    chunks.extend(lc_documents)
  return chunks

chunks = create_chunks(documents)

vector_store.add_documents(chunks)
